from pathlib import Path
import os
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A3,landscape
from reportlab.pdfbase.pdfmetrics import stringWidth
out=Path(os.environ.get('RESCUE_ASSIGNMENT_OUT','submission-output/Progress-Update'));w=load_workbook(out/'Rescue-Reports-Gantt-October-12-30.xlsx',data_only=True);s=w.worksheets[0]
W,H=landscape(A3);c=canvas.Canvas(str(out/'Rescue-Reports-Gantt-October-12-30.pdf'),pagesize=(W,H));c.setTitle('Rescue Reports | Gantt Chart | October 12–30, 2026')
def text(x,y,t,size=10,color='#241747',bold=False):c.setFillColor(HexColor(color));c.setFont('Helvetica-Bold' if bold else 'Helvetica',size);c.drawString(x,y,t)
def wrap(t,x,y,width,size=9,leading=12):
 words=str(t).split();line=''
 for word in words:
  candidate=(line+' '+word).strip()
  if stringWidth(candidate,'Helvetica',size)>width:
   text(x,y,line,size);y-=leading;line=word
  else:line=candidate
 text(x,y,line,size)
text(35,H-42,'RESCUE REPORTS',14,'#6D28D9',True);text(35,H-75,'System Project Gantt Chart',28,bold=True);text(35,H-100,'October 12–30, 2026  |  Updated October 20, 2026',12)
text(35,H-124,'Members: Lopez, Raymond DG; De Leon, Justine; Cabuguang, Liam; [Owner’s submission name to confirm]',10)
text(35,H-146,'V: recorded activity   W: in progress   P: planned   B: reconstructed early baseline, not verified work',10)
left=35;top=H-170;aw=263;tw=221;dw=(W-70-aw-tw)/19;rh=30;hh=39
c.setFillColor(HexColor('#241747'));c.rect(left,top-hh,W-70,hh,fill=1,stroke=0);text(left+8,top-23,'Activity / work package',10,'#FFFFFF',True);text(left+aw+8,top-23,'Target outcome',10,'#FFFFFF',True)
for i in range(19):
 d=s.cell(8,6+i).value
 if not hasattr(d,'strftime'):raise ValueError('Gantt formulas must be recalculated first')
 x=left+aw+tw+i*dw;text(x+5,top-13,d.strftime('%a'),7,'#FFFFFF');text(x+7,top-28,d.strftime('%d'),10,'#FFFFFF',True)
colors={'V':'#DCCDFB','W':'#BDEBDD','P':'#DDE7F7','B':'#E4E4E8'}
for r in range(9,24):
 y=top-hh-(r-8)*rh
 c.setFillColor(HexColor('#F7F5FA' if r%2 else '#FFFFFF'));c.rect(left,y,W-70,rh,fill=1,stroke=0)
 wrap(s.cell(r,1).value,left+7,y+rh-12,aw-14,8.8,10);wrap(s.cell(r,5).value,left+aw+7,y+rh-12,tw-14,8.4,10)
 for i in range(19):
  v=s.cell(r,6+i).value;x=left+aw+tw+i*dw
  if v:c.setFillColor(HexColor(colors[v]));c.roundRect(x+2,y+4,dw-4,rh-8,3,fill=1,stroke=0);text(x+dw/2-3,y+10,v,9,bold=True)
 c.setStrokeColor(HexColor('#E5E1ED'));c.line(left,y,W-35,y)
y=top-hh-15*rh-23
wrap('Basis: verified bars show dated activity windows, not completion of an entire phase or individual work hours. October 12–15 baseline dates are reconstructed planning; precise execution is not established. Future targets depend on available accounts, hardware and owner decisions.',35,y,W-70,9.5,13)
wrap('Interim chart and progress report: target October 21 at 1 PM, based on “tomorrow” in the assignment supplied October 20. Final development window ends October 30. The accompanying progress report excludes October 20 accomplishments.',35,y-44,W-70,9.5,13)
text(35,29,'Editable original: Rescue-Reports-Gantt-October-12-30.xlsx  |  Sources and assumptions tab included',8.5)
c.showPage();c.save()
