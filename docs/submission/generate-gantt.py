from pathlib import Path
from datetime import datetime,timedelta
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
from openpyxl.worksheet.page import PageMargins
import os
out=Path(os.environ.get('RESCUE_ASSIGNMENT_OUT','submission-output/Progress-Update'));out.mkdir(parents=True,exist_ok=True)
w=Workbook();s=w.active;s.title='October 12–30 Gantt'
s.merge_cells('A1:X1');s['A1']='RESCUE REPORTS  |  SYSTEM PROJECT GANTT CHART';s.merge_cells('A2:X2');s['A2']='Development window: October 12–30, 2026  •  Updated October 20, 2026'
s.merge_cells('A3:X3');s['A3']='Members: Lopez, Raymond DG; De Leon, Justine; Cabuguang, Liam; [Owner’s submission name to confirm]'
s.merge_cells('A4:X4');s['A4']='V = activity recorded in dated evidence; W = work in progress; P = planned; B = reconstructed baseline plan (not verified activity).'
s.merge_cells('A5:X5');s['A5']='Editable: activity, status and Start/End dates. Daily cells recalculate from those dates. No completion percentage is asserted.'
s.merge_cells('A6:X6');s['A6']='Interim Gantt + progress submission: target October 21, 1 PM (relative deadline supplied October 20). Final project timeline ends October 30.'
headers=['Activity / work package','Status','Start','End','Target outcome']
for col,h in enumerate(headers,1):s.cell(8,col,h)
for k in range(19):
 c=s.cell(8,6+k);c.value=datetime(2026,10,12) if k==0 else f'={s.cell(8,5+k).coordinate}+1';c.number_format='ddd\ndd'
rows=[
('Requirements, scope and architecture','Baseline',12,15,'Scope and workflow agreed','X-SPEC.md; baseline start dates reconstructed, not dated accomplishments.'),
('Initial schema, auth and app foundation','Baseline',13,16,'Data and identity foundation','X-SPEC.md phases; baseline start dates reconstructed.'),
('Resident and admin web modules','Verified',16,16,'Resident forms, tracking and admin','Git October16: 47252c3, 6439330, 52e8156.'),
('Web UI, motion and page audits','Verified',16,19,'Coherent pages and interaction checks','Git e9c6426; October18 audit commits; October19 release ledger.'),
('API retry safety and cleanup recovery','Verified',18,19,'Stable retries and durable recovery','October18–19 integration and full-repair records; review repairs recorded October19.'),
('Role permissions and live access sync','Verified',18,19,'Role/team changes protect views','October19 RELEASE-LEDGER and native follow-up.'),
('Flutter progress history and retry fixes','Verified',19,19,'Resident/responder hardening','NATIVE-FOLLOWUP-2026-10-19; NIR-01/NIR-02 and 210-test snapshot.'),
('Flutter visual redesign and Android demo','Verified',20,20,'Reviewed theme and test APK','October20 only: 1cf0f33 and FINAL-VERIFICATION; excluded from October12–19 report.'),
('Regression and responsive verification','Verified',18,20,'Local regression evidence','Dated audit records; verified windows are recorded activity, not a claim of all-day work.'),
('Guide, narrated video and delivery package','In progress',20,21,'Reviewed submission materials','Owner request October20; local delivery preparation.'),
('Interim Gantt and progress report','In progress',20,21,'Upload one chart and one report','Assignment screenshot supplied October20, due tomorrow 1 PM.'),
('Provider/device checks and policy inputs','Planned',21,26,'External gates checked where available','Proposed plan, depends on accounts/hardware/owner decisions; not promised complete.'),
('Fix remaining findings and retest','Planned',23,27,'Prioritized defects resolved','Proposed next-phase schedule.'),
('Member review and presentation rehearsal','Planned',26,29,'Confirmed contributions and explanation','Proposed plan; no individual work attributed without confirmation.'),
('Final freeze, archive and handover','Planned',30,30,'Final source and demonstration ready','Proposed target within the required October12–30 window.')]
for r,(activity,status,a,b,target,evidence) in enumerate(rows,9):
 for c,v in enumerate([activity,status,datetime(2026,10,a),datetime(2026,10,b),target],1):s.cell(r,c,v)
 s.cell(r,1).comment=Comment(evidence,'Source / assumption');s.cell(r,3).comment=Comment(evidence,'Source / assumption');s.cell(r,4).comment=Comment(evidence,'Source / assumption')
 for c in [3,4]:s.cell(r,c).number_format='mmm dd'
 for c in range(6,25):s.cell(r,c,f'=IF(AND({s.cell(8,c).coordinate}>=$C{r},{s.cell(8,c).coordinate}<=$D{r}),IF($B{r}="Verified","V",IF($B{r}="In progress","W",IF($B{r}="Baseline","B","P"))),"")')
 s.row_dimensions[r].height=34
for row in s:
 for c in row:
  c.font=Font(name='Arial',size=10,color='241747');c.alignment=Alignment(vertical='center',wrap_text=True)
  if c.row>=8:c.border=Border(bottom=Side(style='hair',color='D8D4E4'))
  if c.column>=6 and c.row>=8:c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
for r in [1,2,8]:
 for c in s[r]:c.fill=PatternFill('solid',fgColor='241747');c.font=Font(name='Arial',size=14 if r==1 else 10,color='FFFFFF',bold=True)
s.row_dimensions[1].height=32;s.row_dimensions[2].height=25;s.row_dimensions[3].height=26;s.row_dimensions[4].height=25;s.row_dimensions[5].height=25;s.row_dimensions[6].height=28;s.row_dimensions[8].height=35
for code,color in [('V','DCCDFB'),('W','BDEBDD'),('P','DDE7F7'),('B','E4E4E8')]:s.conditional_formatting.add('F9:X23',CellIsRule(operator='equal',formula=['"'+code+'"'],fill=PatternFill('solid',fgColor=color)))
for col,width in [('A',36),('B',13),('C',12),('D',12),('E',33)]:s.column_dimensions[col].width=width
for c in range(6,25):s.column_dimensions[s.cell(8,c).column_letter].width=5
s.freeze_panes='F9';s.sheet_view.showGridLines=False;s.print_options.horizontalCentered=True;s.sheet_properties.pageSetUpPr.fitToPage=True
s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1;s.print_area='A1:X23';s.page_margins=PageMargins(left=.3,right=.3,top=.3,bottom=.3)
s.oddFooter.center.text='Rescue Reports | October 12–30, 2026';s.oddFooter.right.text='Page &P'
# A visible evidence/assumptions sheet keeps baseline dates distinct from evidence.
t=w.create_sheet('Sources and assumptions');t.append(['Basis','Detail']);t.append(['Reporting cutoff','Progress report stops October19. Gantt updated October20.']);t.append(['Verified bars','Dated activity windows, not assertions of entire phase completion or member hours.']);t.append(['Baseline bars','October12–15 planning reconstructed to cover assigned timeline; exact daily execution is not established.']);t.append(['Future bars','October21–30 targets are proposed and may depend on external accounts/hardware.']);t.append(['Member names','Three names supplied by owner; fourth full submission name not supplied.']);t.append(['Individual roles','Unconfirmed. No role/contribution assigned to a person by this workbook.'])
for row in rows:t.append([row[0],row[5]])
for row in t:
 for c in row:c.font=Font(name='Arial',size=11);c.alignment=Alignment(wrap_text=True,vertical='top')
t.column_dimensions['A'].width=40;t.column_dimensions['B'].width=110
for r in range(1,t.max_row+1):t.row_dimensions[r].height=35
w.save(out/'Rescue-Reports-Gantt-October-12-30.xlsx')
print(out)
